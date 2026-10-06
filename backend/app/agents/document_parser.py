import pymupdf
import pdfplumber
import docx
import csv
import io
import re
from typing import Optional
from openai import OpenAI
from app.config import settings
from app.agents.extraction_schemas import ParsedAssignment, ParsedRoster, ParsedPolicy
from app.agents.prompts import ASSIGNMENT_PARSE_PROMPT, ROSTER_PARSE_PROMPT
from app.agents.safety import sanitize_document_text

client = None
def get_openai_client():
    global client
    if client is None and settings.OPENAI_API_KEY:
        client = OpenAI(api_key=settings.OPENAI_API_KEY)
    return client

def extract_text_from_pdf(file_path: str) -> str:
    doc = pymupdf.open(file_path)
    text = "\n".join([page.get_text() for page in doc])
    if len(text.strip()) < 20 and len(doc) > 0:
        # Scanned PDF with no text layer: render first page and let the vision LLM read it
        pix = doc[0].get_pixmap(dpi=150)
        import base64 as _b64
        encoded = _b64.b64encode(pix.tobytes("png")).decode("utf-8")
        doc.close()
        return f"[IMAGE:image/png;base64,{encoded}]"
    doc.close()
    return text

def extract_text_from_docx(file_path: str) -> str:
    doc = docx.Document(file_path)
    parts = [p.text for p in doc.paragraphs if p.text]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)

def extract_text_from_csv(file_path: str) -> str:
    with open(file_path, 'r', encoding='utf-8') as f:
        return f.read()

import base64

def extract_text_from_image(file_path: str, mime_type: str) -> str:
    with open(file_path, "rb") as image_file:
        encoded_string = base64.b64encode(image_file.read()).decode('utf-8')
    return f"[IMAGE:{mime_type};base64,{encoded_string}]"

def extract_text(file_path: str, mime_type: str) -> str:
    lower = file_path.lower()
    if lower.endswith(('.jpg', '.jpeg', '.png', '.webp')) and 'image' not in mime_type:
        mime_type = 'image/jpeg' if lower.endswith(('.jpg', '.jpeg')) else f"image/{lower.rsplit('.', 1)[1]}"
    if 'pdf' in mime_type or lower.endswith('.pdf'):
        return extract_text_from_pdf(file_path)
    elif 'docx' in mime_type or 'word' in mime_type or lower.endswith('.docx'):
        return extract_text_from_docx(file_path)
    elif 'csv' in mime_type or lower.endswith('.csv'):
        return extract_text_from_csv(file_path)
    elif 'image' in mime_type:
        return extract_text_from_image(file_path, mime_type)
    elif 'text' in mime_type:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()
    else:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            return f.read()

from app.agents.llm_factory import LLMFactory

def _fallback_due_date(text: str) -> Optional[str]:
    """Regex fallback for explicit dates like 25/10/2026, 2026-10-25, 25 Oct 2026."""
    from datetime import datetime
    m = re.search(r"\b(\d{4})-(\d{1,2})-(\d{1,2})\b", text)
    if m:
        try:
            return datetime(int(m[1]), int(m[2]), int(m[3])).strftime("%Y-%m-%d")
        except ValueError:
            pass
    m = re.search(r"\b(\d{1,2})[/\-.](\d{1,2})[/\-.](\d{2,4})\b", text)
    if m:
        y = int(m[3]) + (2000 if int(m[3]) < 100 else 0)
        try:
            return datetime(y, int(m[2]), int(m[1])).strftime("%Y-%m-%d")  # day-first (India)
        except ValueError:
            pass
    m = re.search(r"\b(\d{1,2})(?:st|nd|rd|th)?\s+([A-Za-z]{3,9})\.?,?\s+(\d{4})\b", text)
    if m:
        for fmt in ("%d %B %Y", "%d %b %Y"):
            try:
                return datetime.strptime(f"{m[1]} {m[2]} {m[3]}", fmt).strftime("%Y-%m-%d")
            except ValueError:
                pass
    return None

async def parse_assignment_document(text: str) -> ParsedAssignment:
    from datetime import datetime, timezone, timedelta
    # India time for "today" so relative dates ("tomorrow") resolve correctly
    now_ist = datetime.now(timezone.utc) + timedelta(hours=5, minutes=30)
    today_str = now_ist.strftime("%Y-%m-%d (%A)")

    image_b64 = None
    mime_type = "image/png"
    raw_text = text

    # Detect image BEFORE sanitizing; sanitize wraps text in tags which hid the marker.
    if text.startswith("[IMAGE:"):
        end_idx = text.find("]")
        if end_idx != -1:
            meta = text[7:end_idx]
            mime_type, image_b64 = meta.split(";base64,")
            raw_text = ""
            prompt = f"Today's Date: {today_str}\n\nExtract assignment information from this image."
    if image_b64 is None:
        safe_text = sanitize_document_text(text)
        prompt = f"Today's Date: {today_str}\n\nExtract assignment information:\n\n{safe_text}"

    try:
        result = LLMFactory.parse_structured(
            prompt=prompt,
            response_schema=ParsedAssignment,
            system_prompt=ASSIGNMENT_PARSE_PROMPT,
            image_b64=image_b64,
            mime_type=mime_type
        )
        if not result.due_date and raw_text:
            result.due_date = _fallback_due_date(raw_text)
        if not result.title:
            if result.subject and result.instructions:
                result.title = f"{result.subject} assignment"
            elif result.instructions:
                words = [w for w in result.instructions.split() if len(w) > 2][:4]
                result.title = " ".join(words).capitalize() if words else "New Assignment"
            else:
                result.ambiguities.append("Title could not be determined")
        if not result.due_date:
            result.ambiguities.append("Due date not found or unclear")
        # Class is chosen at upload time; the caller fills it in, so it's not flagged here.
        return result
    except Exception as e:
        print(f"[DocumentParser] Parsing failed: {e}")
        return ParsedAssignment(
            ambiguities=[f"Document parsing failed: {str(e)}"],
            confidence=0.0
        )

async def parse_roster_document(text: str) -> ParsedRoster:
    safe_text = sanitize_document_text(text)
    try:
        return LLMFactory.parse_structured(
            prompt=f"Extract roster information:\n\n{safe_text}",
            response_schema=ParsedRoster,
            system_prompt=ROSTER_PARSE_PROMPT
        )
    except Exception as e:
        print(f"[DocumentParser] Roster parsing failed: {e}")
        return parse_csv_roster(text)

def parse_csv_roster(text: str) -> ParsedRoster:
    """Fallback CSV parser without LLM."""
    rows = []
    reader = csv.DictReader(io.StringIO(text))
    seen_names = set()
    duplicates = []
    
    for i, row in enumerate(reader):
        # Try common column name variations
        name = row.get('student_name') or row.get('name') or row.get('Student Name') or row.get('Name') or ''
        grade = row.get('grade_class') or row.get('class') or row.get('Grade') or row.get('Class') or ''
        parent = row.get('parent_name') or row.get('Parent Name') or row.get('parent') or ''
        contact = row.get('parent_contact') or row.get('Parent Contact') or row.get('phone') or row.get('Phone') or ''
        notes = row.get('notes') or row.get('Notes') or ''
        
        flags = []
        if not name.strip():
            flags.append('Missing student name')
        if not grade.strip():
            flags.append('Missing class/grade')
        if name.strip().lower() in seen_names:
            flags.append('Possible duplicate')
            duplicates.append({'row': i, 'name': name.strip()})
        seen_names.add(name.strip().lower())
        
        from app.agents.extraction_schemas import ParsedRosterRow
        rows.append(ParsedRosterRow(
            student_name=name.strip(),
            grade_class=grade.strip() or None,
            parent_name=parent.strip() or None,
            parent_contact=contact.strip() or None,
            notes=notes.strip() or None,
            flags=flags
        ))
    
    return ParsedRoster(rows=rows, duplicate_candidates=duplicates)
